"""
Тесты для core/export.py — экспорт в Excel.
"""

import os
import pytest
from openpyxl import load_workbook


# ============================================================
# ЭКСПОРТ ОДНОЙ КАТЕГОРИИ
# ============================================================

class TestExportCategory:
    """Тесты export_category_to_excel."""

    def test_file_created(self, test_db, sample_category, sample_product, export_dir):
        """Файл экспорта создаётся."""
        from core.export import export_category_to_excel

        filepath = export_category_to_excel(sample_category)
        assert filepath is not None
        assert os.path.exists(filepath)

    def test_file_extension(self, test_db, sample_category, sample_product, export_dir):
        """Файл с расширением .xlsx."""
        from core.export import export_category_to_excel

        filepath = export_category_to_excel(sample_category)
        assert filepath.endswith(".xlsx")

    def test_file_in_exports_dir(self, test_db, sample_category, sample_product, export_dir):
        """Файл лежит в папке exports."""
        from core.export import export_category_to_excel

        filepath = export_category_to_excel(sample_category)
        assert str(export_dir) in filepath

    def test_category_not_exists(self, test_db, export_dir):
        """Несуществующая категория — None."""
        from core.export import export_category_to_excel

        result = export_category_to_excel(9999)
        assert result is None

    def test_content_correct(self, test_db, sample_category, sample_product, export_dir):
        """Данные в Excel совпадают с БД."""
        from core.export import export_category_to_excel

        filepath = export_category_to_excel(sample_category)

        wb = load_workbook(filepath)
        ws = wb.active

        found = False
        for row in ws.iter_rows(values_only=True):
            if row and "Компьютер" in str(row):
                found = True
                break
        assert found, "Товар 'Компьютер' не найден в Excel"

    def test_totals_correct(self, test_db, sample_category, export_dir):
        """Итоги правильные."""
        from core.export import export_category_to_excel

        test_db.add_product(sample_category, "Мышь", 10, 1500)
        test_db.add_product(sample_category, "Монитор", 3, 15000)

        filepath = export_category_to_excel(sample_category)
        wb = load_workbook(filepath)
        ws = wb.active

        found_total = False
        for row in ws.iter_rows(values_only=True):
            if row and row[1] == "ИТОГО:":
                assert row[2] == 13
                assert row[4] == 60000
                found_total = True
                break
        assert found_total, "Строка 'ИТОГО' не найдена"

    def test_no_stock_excluded(self, test_db, sample_category, export_dir):
        """Товары с no_stock=True не учитываются."""
        from core.export import export_category_to_excel

        test_db.add_product(sample_category, "Мышь", 10, 1500)
        test_db.add_product(sample_category, "Клавиатура", 0, 2000, no_stock=True)

        filepath = export_category_to_excel(sample_category)
        wb = load_workbook(filepath)
        ws = wb.active

        for row in ws.iter_rows(values_only=True):
            if row and row[1] == "ИТОГО:":
                assert row[2] == 10
                assert row[4] == 15000
                break


# ============================================================
# ЭКСПОРТ ВСЕГО СКЛАДА
# ============================================================

class TestExportAll:
    """Тесты export_all_to_excel."""

    def test_file_created(self, test_db, export_dir):
        """Файл создаётся."""
        from core.export import export_all_to_excel

        test_db.add_category("Продукты")

        filepath = export_all_to_excel()
        assert filepath is not None
        assert os.path.exists(filepath)

    def test_multiple_sheets(self, test_db, export_dir):
        """Файл содержит несколько листов."""
        from core.export import export_all_to_excel

        cat1 = test_db.add_category("Продукты")
        cat2 = test_db.add_category("Электроника")
        test_db.add_product(cat1, "Хлеб", 10, 50)
        test_db.add_product(cat2, "Мышь", 5, 1500)

        filepath = export_all_to_excel()
        wb = load_workbook(filepath)

        sheet_names = wb.sheetnames
        assert "Сводка" in sheet_names
        assert "Продукты" in sheet_names
        assert "Электроника" in sheet_names

    def test_summary_has_totals(self, test_db, export_dir):
        """В листе 'Сводка' есть строка ИТОГО."""
        from core.export import export_all_to_excel

        cat = test_db.add_category("Продукты")
        test_db.add_product(cat, "Хлеб", 10, 50)

        filepath = export_all_to_excel()
        wb = load_workbook(filepath)
        ws = wb["Сводка"]

        found_total = False
        for row in ws.iter_rows(values_only=True):
            if row and row[0] == "ИТОГО:":
                found_total = True
                break
        assert found_total