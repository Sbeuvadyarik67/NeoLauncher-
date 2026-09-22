"""
NeoTracker — экспорт данных в Excel.

Содержит:
- export_category_to_excel() — экспорт товаров одной категории
- export_all_to_excel() — экспорт всего склада (все категории)
"""

import os
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

import database as db


# ============================================================
# СТИЛИ
# ============================================================

HEADER_FONT = Font(bold=True, color="FFFFFF", size=12)
HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
HEADER_ALIGN = Alignment(horizontal="center", vertical="center")

TITLE_FONT = Font(bold=True, size=14)
SUBTITLE_FONT = Font(italic=True, size=10, color="666666")

TOTAL_FONT = Font(bold=True, size=12)
TOTAL_FILL = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")

CELL_ALIGN_CENTER = Alignment(horizontal="center", vertical="center")
CELL_ALIGN_RIGHT = Alignment(horizontal="right", vertical="center")

THIN_BORDER = Border(
    left=Side(style="thin", color="CCCCCC"),
    right=Side(style="thin", color="CCCCCC"),
    top=Side(style="thin", color="CCCCCC"),
    bottom=Side(style="thin", color="CCCCCC"),
)


# ============================================================
# ВСПОМОГАТЕЛЬНОЕ
# ============================================================

def _get_export_dir():
    """Папка для экспорта — рядом с программой, подпапка exports/."""
    base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    export_dir = os.path.join(base_dir, "exports")
    os.makedirs(export_dir, exist_ok=True)
    return export_dir


def _make_filename(prefix):
    """Формирует имя файла с датой и временем."""
    timestamp = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")
    safe_prefix = "".join(c if c.isalnum() or c in "-_" else "_" for c in prefix)
    return f"{safe_prefix}_{timestamp}.xlsx"


def _autosize_columns(ws, min_width=10, max_width=40):
    """Автоматически подгоняет ширину колонок."""
    for col_idx, column in enumerate(ws.columns, start=1):
        max_length = 0
        for cell in column:
            try:
                if cell.value:
                    val = str(cell.value)
                    max_length = max(max_length, len(val) * 1.1)
            except Exception:
                pass
        width = min(max_width, max(min_width, int(max_length) + 2))
        ws.column_dimensions[get_column_letter(col_idx)].width = width


def _style_header_row(ws, row_idx, col_count):
    """Оформляет строку заголовка."""
    for col in range(1, col_count + 1):
        cell = ws.cell(row=row_idx, column=col)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = HEADER_ALIGN
        cell.border = THIN_BORDER


# ============================================================
# ЭКСПОРТ ОДНОЙ КАТЕГОРИИ
# ============================================================

def export_category_to_excel(category_id):
    """
    Экспортирует товары одной категории в Excel.
    Возвращает путь к файлу или None, если категории нет.
    """
    category = db.get_category(category_id)
    if not category:
        return None

    products = db.get_products_by_category(category_id)

    wb = Workbook()
    ws = wb.active
    ws.title = category["name"][:30]

    # ---- Заголовок листа ----
    ws["A1"] = f"Категория: {category['name']}"
    ws["A1"].font = TITLE_FONT
    ws.merge_cells("A1:F1")

    ws["A2"] = f"Дата экспорта: {datetime.now().strftime('%d.%m.%Y %H:%M')}"
    ws["A2"].font = SUBTITLE_FONT
    ws.merge_cells("A2:F2")

    # ---- Шапка таблицы ----
    headers = ["№", "Название", "Количество", "Цена (₽)", "Сумма (₽)", "Описание"]
    header_row = 4
    for col, header in enumerate(headers, start=1):
        ws.cell(row=header_row, column=col, value=header)

    _style_header_row(ws, header_row, len(headers))

    # ---- Данные ----
    total_qty = 0
    total_sum = 0.0
    row_idx = header_row + 1

    for i, p in enumerate(products, start=1):
        if p["no_stock"]:
            qty = 0
            price = p["price"]
            sum_val = 0
            status = "нет в наличии"
        else:
            qty = p["quantity"]
            price = p["price"]
            sum_val = qty * price
            status = ""

        ws.cell(row=row_idx, column=1, value=i).alignment = CELL_ALIGN_CENTER
        ws.cell(row=row_idx, column=2, value=p["name"])
        ws.cell(row=row_idx, column=3, value=qty).alignment = CELL_ALIGN_CENTER
        ws.cell(row=row_idx, column=4, value=price).alignment = CELL_ALIGN_RIGHT
        ws.cell(row=row_idx, column=5, value=sum_val).alignment = CELL_ALIGN_RIGHT
        ws.cell(row=row_idx, column=6, value=p.get("description", "") or status)

        ws.cell(row=row_idx, column=4).number_format = '#,##0.00 ₽'
        ws.cell(row=row_idx, column=5).number_format = '#,##0.00 ₽'

        for col in range(1, len(headers) + 1):
            ws.cell(row=row_idx, column=col).border = THIN_BORDER

        total_qty += qty
        total_sum += sum_val
        row_idx += 1

    # ---- Итоги ----
    total_row = row_idx
    ws.cell(row=total_row, column=1, value="").fill = TOTAL_FILL
    ws.cell(row=total_row, column=2, value="ИТОГО:").font = TOTAL_FONT
    ws.cell(row=total_row, column=2).fill = TOTAL_FILL
    ws.cell(row=total_row, column=3, value=total_qty).font = TOTAL_FONT
    ws.cell(row=total_row, column=3).fill = TOTAL_FILL
    ws.cell(row=total_row, column=3).alignment = CELL_ALIGN_CENTER
    ws.cell(row=total_row, column=4, value="").fill = TOTAL_FILL
    ws.cell(row=total_row, column=5, value=total_sum).font = TOTAL_FONT
    ws.cell(row=total_row, column=5).fill = TOTAL_FILL
    ws.cell(row=total_row, column=5).number_format = '#,##0.00 ₽'
    ws.cell(row=total_row, column=5).alignment = CELL_ALIGN_RIGHT
    ws.cell(row=total_row, column=6, value="").fill = TOTAL_FILL

    for col in range(1, len(headers) + 1):
        ws.cell(row=total_row, column=col).border = THIN_BORDER

    _autosize_columns(ws, min_width=8, max_width=40)

    export_dir = _get_export_dir()
    filename = _make_filename(f"NeoTracker_{category['name']}")
    filepath = os.path.join(export_dir, filename)
    wb.save(filepath)

    return filepath


# ============================================================
# ЭКСПОРТ ВСЕГО СКЛАДА
# ============================================================

def export_all_to_excel():
    """
    Экспортирует все категории и товары в один файл (несколько листов).
    Возвращает путь к файлу.
    """
    categories = db.get_all_categories()

    wb = Workbook()
    ws_summary = wb.active
    ws_summary.title = "Сводка"

    # ---- Сводка по всем категориям ----
    ws_summary["A1"] = "NeoTracker — общий отчёт по складу"
    ws_summary["A1"].font = Font(bold=True, size=16)
    ws_summary.merge_cells("A1:D1")

    ws_summary["A2"] = f"Дата экспорта: {datetime.now().strftime('%d.%m.%Y %H:%M')}"
    ws_summary["A2"].font = SUBTITLE_FONT
    ws_summary.merge_cells("A2:D2")

    headers = ["Категория", "Товаров", "Общее кол-во", "Общая стоимость (₽)"]
    header_row = 4
    for col, header in enumerate(headers, start=1):
        ws_summary.cell(row=header_row, column=col, value=header)
    _style_header_row(ws_summary, header_row, len(headers))

    row_idx = header_row + 1
    grand_qty = 0
    grand_val = 0.0

    for cat in categories:
        products = db.get_products_by_category(cat["id"])
        cat_qty = sum(p["quantity"] for p in products if not p["no_stock"])
        cat_val = sum(p["quantity"] * p["price"] for p in products if not p["no_stock"])

        ws_summary.cell(row=row_idx, column=1, value=cat["name"])
        ws_summary.cell(row=row_idx, column=2, value=len(products)).alignment = CELL_ALIGN_CENTER
        ws_summary.cell(row=row_idx, column=3, value=cat_qty).alignment = CELL_ALIGN_CENTER
        ws_summary.cell(row=row_idx, column=4, value=cat_val).alignment = CELL_ALIGN_RIGHT
        ws_summary.cell(row=row_idx, column=4).number_format = '#,##0.00 ₽'

        for col in range(1, 5):
            ws_summary.cell(row=row_idx, column=col).border = THIN_BORDER

        grand_qty += cat_qty
        grand_val += cat_val
        row_idx += 1

    # ---- Итоги ----
    ws_summary.cell(row=row_idx, column=1, value="ИТОГО:").font = TOTAL_FONT
    ws_summary.cell(row=row_idx, column=1).fill = TOTAL_FILL
    ws_summary.cell(row=row_idx, column=2, value="").fill = TOTAL_FILL
    ws_summary.cell(row=row_idx, column=3, value=grand_qty).font = TOTAL_FONT
    ws_summary.cell(row=row_idx, column=3).fill = TOTAL_FILL
    ws_summary.cell(row=row_idx, column=3).alignment = CELL_ALIGN_CENTER
    ws_summary.cell(row=row_idx, column=4, value=grand_val).font = TOTAL_FONT
    ws_summary.cell(row=row_idx, column=4).fill = TOTAL_FILL
    ws_summary.cell(row=row_idx, column=4).number_format = '#,##0.00 ₽'
    ws_summary.cell(row=row_idx, column=4).alignment = CELL_ALIGN_RIGHT

    for col in range(1, 5):
        ws_summary.cell(row=row_idx, column=col).border = THIN_BORDER

    _autosize_columns(ws_summary, min_width=12, max_width=40)

    # ---- Отдельный лист на каждую категорию ----
    for cat in categories:
        products = db.get_products_by_category(cat["id"])
        sheet_name = cat["name"][:30] or f"Кат_{cat['id']}"
        for ch in [":", "\\", "/", "?", "*", "[", "]"]:
            sheet_name = sheet_name.replace(ch, "_")

        ws = wb.create_sheet(title=sheet_name)

        ws["A1"] = f"Категория: {cat['name']}"
        ws["A1"].font = TITLE_FONT
        ws.merge_cells("A1:E1")

        headers = ["№", "Название", "Количество", "Цена (₽)", "Сумма (₽)"]
        header_row = 3
        for col, header in enumerate(headers, start=1):
            ws.cell(row=header_row, column=col, value=header)
        _style_header_row(ws, header_row, len(headers))

        row_idx = header_row + 1
        total_qty = 0
        total_sum = 0.0

        for i, p in enumerate(products, start=1):
            if p["no_stock"]:
                qty = 0
                price = p["price"]
                sum_val = 0
            else:
                qty = p["quantity"]
                price = p["price"]
                sum_val = qty * price

            ws.cell(row=row_idx, column=1, value=i).alignment = CELL_ALIGN_CENTER
            ws.cell(row=row_idx, column=2, value=p["name"])
            ws.cell(row=row_idx, column=3, value=qty).alignment = CELL_ALIGN_CENTER
            ws.cell(row=row_idx, column=4, value=price).alignment = CELL_ALIGN_RIGHT
            ws.cell(row=row_idx, column=4).number_format = '#,##0.00 ₽'
            ws.cell(row=row_idx, column=5, value=sum_val).alignment = CELL_ALIGN_RIGHT
            ws.cell(row=row_idx, column=5).number_format = '#,##0.00 ₽'

            for col in range(1, 6):
                ws.cell(row=row_idx, column=col).border = THIN_BORDER

            total_qty += qty
            total_sum += sum_val
            row_idx += 1

        ws.cell(row=row_idx, column=1, value="").fill = TOTAL_FILL
        ws.cell(row=row_idx, column=2, value="ИТОГО:").font = TOTAL_FONT
        ws.cell(row=row_idx, column=2).fill = TOTAL_FILL
        ws.cell(row=row_idx, column=3, value=total_qty).font = TOTAL_FONT
        ws.cell(row=row_idx, column=3).fill = TOTAL_FILL
        ws.cell(row=row_idx, column=3).alignment = CELL_ALIGN_CENTER
        ws.cell(row=row_idx, column=4, value="").fill = TOTAL_FILL
        ws.cell(row=row_idx, column=5, value=total_sum).font = TOTAL_FONT
        ws.cell(row=row_idx, column=5).fill = TOTAL_FILL
        ws.cell(row=row_idx, column=5).number_format = '#,##0.00 ₽'
        ws.cell(row=row_idx, column=5).alignment = CELL_ALIGN_RIGHT

        for col in range(1, 6):
            ws.cell(row=row_idx, column=col).border = THIN_BORDER

        _autosize_columns(ws, min_width=8, max_width=40)

    export_dir = _get_export_dir()
    filename = _make_filename("NeoTracker_Весь_склад")
    filepath = os.path.join(export_dir, filename)
    wb.save(filepath)

    return filepath